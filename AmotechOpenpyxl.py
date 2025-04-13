import openpyxl
import psutil
from GLforAmotech import *
from MyPPrint import *
from Voucher import *
from Vouchers import *

def LoadExcel(filename, sheetname):
    wb = openpyxl.load_workbook(filename, data_only=True)
    ws = wb[sheetname]
    data = [[cell.value for cell in row] for row in ws.iter_rows()]
    wb.close()
    return data

def SaveExcel(filename, data, sheet_name="Sheet1"):
    try:
        wb = openpyxl.load_workbook(filename)
    except FileNotFoundError:
        wb = openpyxl.Workbook()
    
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(title=sheet_name)
    
    for row in data:
        ws.append(row)
    
    wb.save(filename)
    wb.close()
    return data

def GetExcelData():
    filename = "C:\\DataTest\\아모텍_분개장 (2024_12).xlsx"
    sheetname = "Sheet1"  
    data = LoadExcel(filename, sheetname)
    return data

if __name__ == "__main__":    
    targetFileName = "C:\\Users\\young\\Downloads\\result.xlsx"
    data = GetExcelData()
    vouchers = Vouchers()
    vouchers.getVouchers(data)

    trialNo = 0
    while  1:
        print("다음 중 원하는 작업을 선택하시요")
        print("   1. 합계시산표를 만들기")
        print("   2. 매출이 (-)인 전표를 추출하기")
        print("   3. 매출전표들의 차대변 합계를 구하기")
        print("   4. 차대변 합계가 다른 전표 확인하기")
        print("   5. 전표금액의 범위별 숫자확인하기")
        print("   6. 전표금액의 범위별 전표추출하기")
        print("   9. 종료하기")
        selection = input("선택:  ")
        trialNo += 1
        match selection:
            case '1':
                print(f"\nHello {trialNo}")
                glforAmotech = Amotech(data)
                glforAmotech.getTrialBalance()
                trialBalance = glforAmotech.changeTrialBalanceList()
                SaveExcel(targetFileName, sorted(trialBalance), "TrialBalance")
                glforAmotech.printTrialBalance()
            case '2':
                print(f"\nHello {trialNo}")
                filtered = vouchers.minusSalesTransactions()    
                filteredList = []    
                for voucher in filtered:
                     filteredList.extend(voucher.ToList())
                SaveExcel(targetFileName, filteredList, "MinusSales")
            case '3':
                print(f"\nHello {trialNo}")
                vouchers.salesTransactions()
            case '4':
                print(f"\nHello {trialNo}")
                vouchers.testVoucherAmount()
            case '5':
                print(f"\nHello {trialNo}")           
                result = vouchers.getVouchersAmounts()
                SaveExcel(targetFileName, result, "VoucherAmount")
            case '6':
                print(f"\nHello {trialNo}")
                filtered = vouchers.ExtraOrdinaryTransactions(7500000000, 100000000000)
                filteredList = []
                for voucher in filtered:
                    filteredList.extend(voucher.ToList())
                SaveExcel(targetFileName, filteredList, "ExtraOrdinary")
            case '9':
                break 
            case _ : 
                continue 
