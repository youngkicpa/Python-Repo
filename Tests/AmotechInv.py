import openpyxl
import openpyxl
from openpyxl import Workbook

# 소스 파일 경로
source_file = r"C:\Users\young\다산\우리팀 - 문서\9. 아모텍\7. 재고및금융실사\재고실사 최종자료\(아모텍) 타처 재고 리스트(전체)_yskim.xlsx"
# 결과 파일 경로
result_file = r"C:\Users\young\Downloads\result.xlsx"

# 소스 파일 열기
wb = openpyxl.load_workbook(source_file)
result_wb = Workbook()
result_ws = result_wb.active

# 첫 번째 시트의 제목 행 복사
first_sheet = wb[wb.sheetnames[0]]
result_ws.append([cell.value for cell in first_sheet[3]])

# 각 시트의 데이터를 합치기 (숨겨진 시트는 스킵)
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    if sheet.sheet_state == 'hidden':
        continue
    for row in sheet.iter_rows(min_row=4, min_col=2, values_only=True):
        if row[0] != "계":
            result_ws.append(row)

# 결과 파일 저장
result_wb.save(result_file)
print(f"데이터가 {result_file}에 성공적으로 저장되었습니다.")