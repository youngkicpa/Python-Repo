# %%
import openpyxl as xl
import pandas as pd
from itertools import islice
import psutil
import os

def get_memory_usage():
    process = psutil.Process()
    memory_info = process.memory_info()
    return memory_info.rss  # RSS (Resident Set Size) 메모리 사용량 반환

print(get_memory_usage())
wb = xl.load_workbook('한화엔진_계정별원장_24년 12월.xlsx', read_only=True, data_only=True)
ws = wb.active

print(get_memory_usage())
debit = 0
credit = 0

for x in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=42, max_col=43, values_only=True):
    if x[0] is not None:
        if type(x[0]) == str:
            print(x[0])
            new = x[0].replace(',', '')
            debit += int(new)
        else:
            debit += x[0]
    if x[1] is not None:
        if type(x[1]) == str:
            print(x[1])
            new = x[1].replace(',', '')
            credit += int(new)
        else:
            credit += x[1]
print("debit: {0:,}", debit)
print("credit: {0:,}", credit)

print(get_memory_usage())

# %%

wb.close()

# %%
