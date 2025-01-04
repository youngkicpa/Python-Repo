# %%
import openpyxl as xl
import pandas as pd
from dozonGL import DozonGLDataframe

class MyXL:
    def __init__(self, fileName=None, sheetName=None, readOnly=False):
        if fileName:
            self.workbook = xl.load_workbook(fileName, read_only=readOnly)
            self.fileName = fileName
        else:
            self.workbook = xl.Workbook()
            self.fileName = None

        if sheetName:
            self.currentWorksheet = self.workbook[sheetName]
        else:
            self.currentWorksheet = self.workbook.active

    def save(self, fileName):
        self.workbook.save(fileName)

    def changeSheet(self, sheetName=None):
        if sheetName and sheetName in self.workbook.sheetnames:
            self.currentWorksheet = self.workbook[sheetName]
        elif sheetName and sheetName not in self.workbook.sheetnames:
            self.currentWorksheet = self.workbook.create_sheet(sheetName)
        else:
            self.currentWorksheet = self.workbook.create_sheet()

    def getCell(self, row, col):
        return self.currentWorksheet.cell(row=row, column=col).value

    def getRow(self, row):
        result = []
        for col in self.currentWorksheet[row]:
            result.append(col.value)
        
        return result

    def getRows(self, startRow, endRow):
        result = []    
        for row in self.currentWorksheet.iter_rows(min_row=startRow, max_row=endRow):
            data = []
            for col in row:
                data.append(col.value)
            result.append(data)
        
        return result

    def getRange(self, cellRange):
        result = []
        for row in self.currentWorksheet[cellRange]:
            data = []
            for col in row:
                data.append(col.value)
            result.append(data)
        
        return result

    def getUsedRangeValues(self):
        min_row, min_col, max_row, max_col = self.currentWorksheet.min_row, self.currentWorksheet.min_column, self.currentWorksheet.max_row, self.currentWorksheet.max_column 
        
        # Retrieve values from the used range 
        used_range_values = [] 
        for row in self.currentWorksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col): 
            used_range_values.append([cell.value for cell in row]) 
            
        return used_range_values

    def saveDataFrame(self, df, sheetName=None):
        self.changeSheet(sheetName)
        self._addIndexTitle(df)
        self._addColumnTitles(df)
        self._addDataFrameRows(df)

    def _addIndexTitle(self, df):
        index_num = 1
        index_title = tuple(df.index.names)
        if isinstance(index_title, tuple):
            for x in index_title:
                self.currentWorksheet.cell(row=1, column=index_num, value=x)
                index_num += 1
        else:
            self.currentWorksheet.cell(row=1, column=index_num, value=index_title)
            index_num += 1

    def _addColumnTitles(self, df):
        index_num = len(df.index.names) if isinstance(df.index.names, tuple) else 1
        for col_num, column_title in enumerate(df.columns, index_num + 1):
            self.currentWorksheet.cell(row=1, column=col_num, value=column_title)

    def _addDataFrameRows(self, df):
        for row_num, row_data in enumerate(df.values, 1):
            index_value = df.index[row_num - 1]
            index_num = 1
            if isinstance(index_value, tuple):
                for x in index_value:
                    self.currentWorksheet.cell(row=row_num + 1, column=index_num, value=x)
                    index_num += 1
            else:
                self.currentWorksheet.cell(row=row_num + 1, column=1, value=index_value)
                index_num += 1
            for col_num, cell_value in enumerate(row_data, index_num):
                self.currentWorksheet.cell(row=row_num + 1, column=col_num, value=cell_value)

def printDataFrameInfo(df):
    print("DataFrame Information:")
    print("Columns:", df.columns.tolist())
    print("Number of rows:", len(df))
    print("Number of columns:", len(df.columns))
    print("Index Titiles: ", df.index.names)
    print("First few rows:")
    print(df.head())

if __name__ == "__main__":
    fileName = "C:\\DataTest\\분개장_KSsystem_2409.xlsx"
    myxl = MyXL(fileName)
    #rangeTuple = getUsedRangeValues(ws)
    #print(rangeTuple)
    #데이타프레임이 어떻게 표현되는지 확인 필요함.
    source = myxl.getUsedRangeValues()
    mygl = DozonGLDataframe(source)
    trialBalance = mygl.getTrialBalanceKSsystem()

    tagetFileName = "C:\\Data\\Programming\\python\\Tests\\resultTB.xlsx"
    targetxl = MyXL()
    #printDataFrameInfo(trialBalance)
    targetxl.saveDataFrame(trialBalance, sheetName="New")
    targetxl.save(tagetFileName)
    myxl.save(myxl.fileName)
# %%
